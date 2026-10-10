#!/usr/bin/env python3
"""Lesson 4: lookups.

XLOOKUP is not in Calc 24.2, so every XLOOKUP the lesson prints is computed by
`formulas` (the Excel-compatible calculator lab.sh pins): one cell at a time
through `engine.xl`, and a whole filled column at a time through `column()`
below, which writes the formula into every row of an .xlsx holding the
tables, the way the fill handle would, and reads back the column and its sum.
VLOOKUP, INDEX and MATCH are computed by Calc, and the margin total is
computed both ways, which is the agreement the lesson points at.

`formulas` runs before Calc starts, because the UNO bridge replaces Python's
import machinery in a way `formulas` cannot load under.
"""
import datetime
import os
import re
import tempfile

from engine import Book, quoted, rows, serial, table_cells, xl
from calcref import calc, excelish

L = "le-srknj3h7"
TABLES = {}
for name in ("Sales", "Products", "Customers"):
    TABLES.update(table_cells(name))
for r in range(2, 110):
    TABLES[("Sales", f"H{r}")] = TABLES[("Sales", f"E{r}")] * TABLES[("Sales", f"F{r}")]


def X(formula, extra=None, check=True):
    """One XLOOKUP formula as printed. `xl` computes it in a cell of a sheet
    called Scratch, so a copy of `Sales` is put there too, and the formula's
    own-sheet references (D2, $O$2:$O$4) read the sales as they do in Excel."""
    if check:
        quoted(L, formula)
    cells = dict(TABLES)
    cells.update(extra or {})
    cells.update({("Scratch", a): v for (sh, a), v in list(cells.items()) if sh == "Sales" and a != "A1"})
    return xl(cells, formula)


def column(formulas_by_col, extra=None, sums=()):
    """Fill each formula (written for row 2) down to row 109 in an .xlsx and
    compute it with `formulas`. Returns {(col, row): value} and the sums."""
    import formulas
    import openpyxl
    cells = dict(TABLES)
    cells.update(extra or {})
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for (sh, addr), v in cells.items():
        ws = wb[sh] if sh in wb.sheetnames else wb.create_sheet(sh)
        ws[addr] = serial(v) if isinstance(v, datetime.date) else v
    ws = wb["Sales"]
    for col, f in formulas_by_col.items():
        for r in range(2, 110):
            ws[f"{col}{r}"] = re.sub(r"(?<![$A-Z])([A-Z])2\b", lambda m: f"{m.group(1)}{r}", f)
    for i, (col, _) in enumerate(sums):
        ws[f"Z{i + 1}"] = f"=SUM({col}2:{col}109)"
    d = tempfile.mkdtemp(prefix="xl-")
    path = os.path.join(d, "book.xlsx")
    wb.save(path)
    sol = formulas.ExcelModel().loads(path).finish().calculate()
    out = {}
    for k, v in sol.items():
        m = re.search(r"SALES'!([A-Z]+)(\d+)$", k.upper())
        if m:
            val = v.value[0][0]
            out[(m.group(1), int(m.group(2)))] = val.item() if hasattr(val, "item") else val
    return out


# --- XLOOKUP, with formulas ------------------------------------------------
print("J2", X("=XLOOKUP(D2,Products!$A$2:$A$7,Products!$F$2:$F$7)"))
print("K2", X("=XLOOKUP(D2,Products!$A$2:$A$7,Products!$G$2:$G$7)"))
quoted(L, "=H2-E2*K2")
quoted(L, "=SUM(L2:L109)")
got = column({"K": "=XLOOKUP(D2,Products!$A$2:$A$7,Products!$G$2:$G$7)", "L": "=H2-E2*K2"},
             sums=[("L", "margin")])
print("XLOOKUP cost and margin: L2", got[("L", 2)], "| =SUM(L2:L109)", got[("Z", 1)])
print("not found:", X('=XLOOKUP("CER2K",Products!$A$2:$A$7,Products!$F$2:$F$7,"not found")'))
print("decaf:", X('=XLOOKUP("Decaf 250 g",Products!$B$2:$B$7,Products!$A$2:$A$7)'))
print("name of C2:", X("=XLOOKUP(C2,Customers!$A$2:$A$12,Customers!$B$2:$B$12)"), TABLES[("Sales", "C2")])
print("trailing space:", X('=XLOOKUP("CER1K ",Products!$A$2:$A$7,Products!$F$2:$F$7)'))
print("lower case:", X('=XLOOKUP("cer1k",Products!$A$2:$A$7,Products!$F$2:$F$7)', check=False))
print("C00 city, row 4:", repr(X("=XLOOKUP(C4,Customers!$A$2:$A$12,Customers!$D$2:$D$12)")),
      TABLES[("Sales", "C4")])
print("C00 city &\"\":", repr(X('=XLOOKUP(C4,Customers!$A$2:$A$12,Customers!$D$2:$D$12)&""')))
print("C04 city &\"\" (row 3):", repr(X('=XLOOKUP(C3,Customers!$A$2:$A$12,Customers!$D$2:$D$12)&""',
                                         check=False)), TABLES[("Sales", "C3")])
bands = {("Sales", "O1"): "From", ("Sales", "P1"): "Band", ("Sales", "O2"): 1, ("Sales", "P2"): "Small",
         ("Sales", "O3"): 4, ("Sales", "P3"): "Medium", ("Sales", "O4"): 10, ("Sales", "P4"): "Large"}
quoted(L, "=XLOOKUP(E2,$O$2:$O$4,$P$2:$P$4,,-1)")
got = column({"M": "=XLOOKUP(E2,$O$2:$O$4,$P$2:$P$4,,-1)"}, extra=bands)
seen = [got[("M", r)] for r in range(2, 110)]
print("XLOOKUP bands: M2", got[("M", 2)], "| Large/Medium/Small",
      [seen.count(x) for x in ("Large", "Medium", "Small")])
for r in range(2, 110):
    e = TABLES[("Sales", f"E{r}")]
    want = "Large" if e >= 10 else "Medium" if e >= 4 else "Small"
    assert got[("M", r)] == want, (r, e, got[("M", r)])
prices = {("Sales", "O6"): "From", ("Sales", "P6"): "CER1K", ("Sales", "O7"): datetime.date(2025, 1, 1),
          ("Sales", "O8"): datetime.date(2026, 1, 1), ("Sales", "P7"): 115, ("Sales", "P8"): 118}
print("price on the day, row 4:", X("=XLOOKUP(B4,$O$7:$O$8,$P$7:$P$8,,-1)", prices),
      [TABLES[("Sales", f"{c}4")] for c in "ADFG"],
      "| row 85:", X("=XLOOKUP(B85,$O$7:$O$8,$P$7:$P$8,,-1)", prices, check=False),
      [TABLES[("Sales", f"{c}85")] for c in "ADFG"])
print("exercise: price on the day, row 64:", X("=XLOOKUP(B64,$O$7:$O$8,$P$7:$P$8,,-1)", prices, check=False),
      [TABLES[("Sales", f"{c}64")] for c in "ABDFG"])
# Do online and shop sales pay the list price of their day? (2025 list = today's less 3)
listp = {r[0]: r[5] for r in rows("Products")[1:]}
off = [(r[0], r[5]) for r in rows("Sales")[1:] if r[6] != "Wholesale"
       and r[5] != listp[r[3]] - (3 if r[1].year == 2025 else 0)]
print("non-wholesale sales not at the list price of their day:", off)

# --- Calc -----------------------------------------------------------------
b = excelish(Book())


def ev(f, addr="Z1", check=True):
    b.set("Sales", addr, calc(quoted(L, f) if check else f))
    return b.value("Sales", addr)


def col(f, letter, check=True):
    b.set("Sales", f"{letter}2", calc(quoted(L, f) if check else f))
    b.fill("Sales", f"{letter}2", f"{letter}109")


try:
    b.put("Sales", [["Revenue"]], 0, 7)
    col("=E2*F2", "H", check=False)
    col("=VLOOKUP(D2,Products!$A$2:$G$7,6,FALSE)", "J")
    print("VLOOKUP J2", b.value("Sales", "J2"))
    print("MATCH", ev("=MATCH(D2,Products!$A$2:$A$7,0)"), "| INDEX", ev("=INDEX(Products!$G$2:$G$7,4)"))
    col("=INDEX(Products!$G$2:$G$7,MATCH(D2,Products!$A$2:$A$7,0))", "K")
    col("=H2-E2*K2", "L")
    print("INDEX/MATCH cost K2", b.value("Sales", "K2"), "| margin L2", b.value("Sales", "L2"),
          "| =SUM(L2:L109)", ev("=SUM(L2:L109)"), "| revenue", b.ev("=SUM(H2:H109)", "Sales", "Z2"))
    print("two MATCHes:", ev('=INDEX(Products!$A$1:$G$7,MATCH("MOG250",Products!$A$1:$A$7,0),'
                            'MATCH("Unit cost",Products!$A$1:$G$1,0))'))
    print("exercise: SUL1K Grams by two MATCHes:", ev('=INDEX(Products!$A$1:$G$7,MATCH("SUL1K",Products!$A$1:$A$7,0),'
                                                     'MATCH("Grams",Products!$A$1:$G$1,0))', check=False),
          "| MATCH over A2:A7 inside INDEX over A1:G7:",
          ev('=INDEX(Products!$A$1:$G$7,MATCH("MOG250",Products!$A$2:$A$7,0),MATCH("Unit cost",Products!$A$1:$G$1,0))',
             check=False))
    print("VLOOKUP CER2K:", ev('=VLOOKUP("CER2K",Products!$A$2:$G$7,6,FALSE)'))
    print("=LEN(D2):", ev("=LEN(D2)"), '| =LEN("CER1K "):', ev('=LEN("CER1K ")', check=False))
    col('=IFNA(VLOOKUP(D2,Products!$A$2:$G$7,6,FALSE),"not in Products")', "N")
    print("IFNA column: rows saying not in Products",
          b.ev('=COUNTIF(N2:N109;"not in Products")', "Sales", "Z3"), "| N2", b.value("Sales", "N2"))
    for r in range(2, 110):
        b.sheet("Sales").getCellRangeByName(f"N{r}").setString("")
    print("C00 sales:", ev('=COUNTIF(C2:C109,"C00")', check=False), "| CER1K sales:",
          ev('=COUNTIF(D2:D109,"CER1K")', check=False), "| sales:", ev("=COUNTA(A2:A109)", check=False))
    print("INDEX on the empty city:", repr(ev("=INDEX(Customers!$D$2:$D$12;MATCH(C4;Customers!$A$2:$A$12;0))",
                                               check=False)))
    b.put("Sales", [["From", "Band"], [1, "Small"], [4, "Medium"], [10, "Large"]], 0, 14)
    col("=VLOOKUP(E2,$O$2:$P$4,2,TRUE)", "M")
    print("VLOOKUP TRUE bands: M2", b.value("Sales", "M2"), "| Large/Medium/Small",
          [b.ev(f'=COUNTIF(M2:M109;"{x}")', "Sales", "Z4") for x in ("Large", "Medium", "Small")])
    # Trap 1: a column inserted in Products before List price.
    b.sheet("Products").Columns.insertByIndex(5, 1)
    print("after inserting a column at F in Products: J2 (VLOOKUP, 6)", b.value("Sales", "J2"),
          "| =SUM(J2:J109)", b.ev("=SUM(J2:J109)", "Sales", "Z5"), "| K2 (INDEX) still", b.value("Sales", "K2"),
          b.sheet("Sales").getCellRangeByName("K2").getFormula())
finally:
    b.close()
