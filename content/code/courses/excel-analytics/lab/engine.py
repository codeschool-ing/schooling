"""A spreadsheet that runs the course's formulas: LibreOffice Calc, driven
through its UNO bridge, holding the data lesson 1 gives the student.

WHY CALC AND NOT EXCEL. The machine the course was written on has no Excel
and cannot have one. Calc implements the same functions under the same names
with the same arguments for everything lessons 2 to 12 compute, so a formula
the lesson prints is put into a cell here exactly as printed and the number
the lesson quotes is what came back. The two places that is not true are
named where they happen: XLOOKUP and XMATCH arrived in Calc 24.8 and this is
24.2, so those go to `formulas` (an Excel-compatible calculator in Python,
pinned in lab.sh) through `xl()`; and Power Query and DAX have no engine
outside Excel at all, so lessons 13 to 16 compute what those steps produce
in plain Python, and their prose says nothing was run in Excel.

THE DATA IS READ OUT OF LESSON 1. `data_blocks()` takes the three tab-separated
fences of section `your-data`, so a number in any lesson is computed from the
bytes the student pastes. A value is typed the way Excel types a pasted one:
a whole number becomes a number, `yyyy-mm-dd` becomes a date, anything else
stays text, and an empty field is an empty cell.

THE FORMULA IS READ OUT OF THE LESSON TOO. `quoted()` refuses a formula that
the lesson's English prose does not contain byte for byte, so the formula the
capture computed and the formula the student types cannot drift apart.
"""
import datetime
import glob
import os
import re
import subprocess
import tempfile
import time

HERE = os.path.dirname(os.path.abspath(__file__))
COURSE = os.path.dirname(HERE)
LESSONS = os.path.join(COURSE, "lessons")
LESSON1 = os.path.join(LESSONS, "le-drcqf3q5")

ISO = re.compile(r"^\d{4}-\d{2}-\d{2}$")
WHOLE = re.compile(r"^-?\d+$")
EPOCH = datetime.date(1899, 12, 30)


def data_blocks():
    """The three tables of lesson 1 section `your-data`, as their text."""
    text = open(os.path.join(LESSON1, "your-data.md"), encoding="utf-8").read()
    out = {}
    for body in re.findall(r"```\n(.*?)\n```", text, re.S):
        first = body.split("\n", 1)[0].split("\t")[0]
        name = {"Sale": "Sales", "Code": "Products", "Customer": "Customers"}.get(first)
        if name:
            out[name] = body
    return out


def parse(field):
    if field == "":
        return None
    if WHOLE.match(field):
        return int(field)
    if ISO.match(field):
        return datetime.date.fromisoformat(field)
    return field


def rows(name):
    """A table as lists of typed values, header first."""
    lines = data_blocks()[name].split("\n")
    return [lines[0].split("\t")] + [[parse(f) for f in l.split("\t")] for l in lines[1:]]


def serial(d):
    return (d - EPOCH).days


def date_of(n):
    return EPOCH + datetime.timedelta(days=int(n))


def quoted(lesson, formula):
    """The formula as the lesson prints it, or a refusal."""
    d = os.path.join(LESSONS, lesson)
    for path in glob.glob(os.path.join(d, "*.md")):
        if path.endswith(".pt.md"):
            continue
        if formula in open(path, encoding="utf-8").read():
            return formula
    raise SystemExit(f"{lesson}: the lesson does not print {formula!r}")


def api(formula):
    """Excel's English spelling to Calc's API grammar: `;` between arguments.
    Commas inside a string literal stay commas."""
    out, q = [], False
    for ch in formula:
        if ch == '"':
            q = not q
        out.append(";" if ch == "," and not q else ch)
    return "".join(out)


class Book:
    """One Calc document with the three tables pasted in, A1 down."""

    def __init__(self, sheets=("Sales", "Products", "Customers"), extra=None):
        import uno
        from com.sun.star.beans import PropertyValue
        self.uno = uno
        prof = tempfile.mkdtemp(prefix="xlab-")
        pipe = "xlab" + str(os.getpid())
        self.proc = subprocess.Popen(
            ["soffice", "--headless", "--invisible", "--norestore", "--nologo",
             f"-env:UserInstallation=file://{prof}", f"--accept=pipe,name={pipe};urp;"],
            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        local = uno.getComponentContext()
        res = local.ServiceManager.createInstanceWithContext("com.sun.star.bridge.UnoUrlResolver", local)
        for _ in range(120):
            try:
                ctx = res.resolve(f"uno:pipe,name={pipe};urp;StarOffice.ComponentContext")
                break
            except Exception:
                time.sleep(0.5)
        else:
            raise SystemExit("LibreOffice did not start")
        self.desk = ctx.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", ctx)
        pv = PropertyValue()
        pv.Name, pv.Value = "Hidden", True
        self.doc = self.desk.loadComponentFromURL("private:factory/scalc", "_blank", 0, (pv,))
        tables = {n: rows(n) for n in sheets}
        tables.update(extra or {})
        first = True
        for name, data in tables.items():
            if first:
                self.doc.Sheets.getByIndex(0).setName(name)
                first = False
            else:
                self.doc.Sheets.insertNewByName(name, self.doc.Sheets.getCount())
            self.put(name, data)
        self.doc.Sheets.insertNewByName("Scratch", self.doc.Sheets.getCount())

    def sheet(self, name):
        return self.doc.Sheets.getByName(name)

    def put(self, name, data, top=0, left=0):
        sh = self.sheet(name)
        for r, row in enumerate(data):
            for c, v in enumerate(row):
                cell = sh.getCellByPosition(left + c, top + r)
                if v is None:
                    continue
                if isinstance(v, datetime.date):
                    cell.setValue(serial(v))
                elif isinstance(v, (int, float)) and not isinstance(v, bool):
                    cell.setValue(v)
                else:
                    cell.setString(str(v))

    def set(self, sheet, addr, formula):
        self.sheet(sheet).getCellRangeByName(addr).setFormula(api(formula))

    def fill(self, sheet, first, last):
        """Copy the formula in `first` down to `last`, as dragging the fill
        handle does: relative references move, `$` ones do not."""
        sh = self.sheet(sheet)
        src = sh.getCellRangeByName(first).getRangeAddress()
        col = re.match(r"[A-Z]+", first).group(0)
        a, b = int(first[len(col):]), int(last[len(col):])
        for r in range(a + 1, b + 1):
            dest = sh.getCellRangeByName(f"{col}{r}").getCellAddress()
            sh.copyRange(dest, src)

    def value(self, sheet, addr):
        self.doc.calculateAll()
        cell = self.sheet(sheet).getCellRangeByName(addr)
        if cell.getError():
            return cell.getString()
        t = cell.getType().value
        if t == "TEXT":
            return cell.getString()
        if t == "FORMULA":
            rt = cell.FormulaResultType2
            if rt == 2:  # TEXT
                return cell.getString()
            if rt == 0:
                return ""
            v = cell.getValue()
            return int(v) if v == int(v) else v
        if t == "EMPTY":
            return ""
        v = cell.getValue()
        return int(v) if v == int(v) else v

    def ev(self, formula, sheet="Scratch", addr="A1"):
        self.set(sheet, addr, formula)
        return self.value(sheet, addr)

    def close(self):
        try:
            self.doc.close(True)
        except Exception:
            pass
        try:
            self.desk.terminate()
        except Exception:
            pass
        try:
            self.proc.wait(timeout=30)
        except Exception:
            self.proc.kill()


def xl(cells, formula):
    """XLOOKUP and XMATCH, which this Calc lacks: the formula is put in an
    .xlsx beside the cells it needs and computed by `formulas`. `cells` maps
    (sheet, address) to a value."""
    import formulas
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for (sh, addr), v in cells.items():
        ws = wb[sh] if sh in wb.sheetnames else wb.create_sheet(sh)
        ws[addr] = serial(v) if isinstance(v, datetime.date) else v
    ws = wb["Scratch"] if "Scratch" in wb.sheetnames else wb.create_sheet("Scratch")
    ws["A1"] = formula
    d = tempfile.mkdtemp(prefix="xl-")
    path = os.path.join(d, "book.xlsx")
    wb.save(path)
    model = formulas.ExcelModel().loads(path).finish()
    sol = model.calculate()
    for k, v in sol.items():
        if k.upper().endswith("SCRATCH'!A1"):
            out = v.value[0][0]
            if hasattr(out, "item"):
                out = out.item()
            if isinstance(out, float) and out == int(out):
                out = int(out)
            return str(out) if not isinstance(out, (int, float, str)) else out
    raise SystemExit("formulas returned nothing for " + formula)


def table_cells(name, sheet=None):
    """A table's values keyed (sheet, address), for `xl()`."""
    out = {}
    for r, row in enumerate(rows(name)):
        for c, v in enumerate(row):
            if v is not None:
                out[(sheet or name, f"{chr(65 + c)}{r + 1}")] = v
    return out
